from __future__ import annotations

from pathlib import Path
from shutil import copy2
from hashlib import sha256
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / 'Template laporan .docx'
BACKUP = ROOT / 'Template laporan - asli.docx'
EVIDENCE = ROOT / 'evidence'
FORMAL = ROOT / '240605110088_Erai_Bagusalim_Mobile_Programming_Practicum_A_Modul_12_Akses_Lokasi_GPS.docx'
SHORT = ROOT / 'Laporan Praktikum Modul 12.docx'

if not BACKUP.exists():
    copy2(TEMPLATE, BACKUP)

doc = Document(BACKUP)

# Helpers

def set_cell_border(cell, **kwargs):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if edge in kwargs:
            edge_el = borders.find(qn(f'w:{edge}'))
            if edge_el is None:
                edge_el = OxmlElement(f'w:{edge}')
                borders.append(edge_el)
            for key, value in kwargs[edge].items():
                edge_el.set(qn(f'w:{key}'), str(value))


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_width(cell, cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn('w:tcW'))
    if tc_w is None:
        tc_w = OxmlElement('w:tcW')
        tc_pr.append(tc_w)
    # Word expects width in twips (1 cm ≈ 567 twips), not EMU.
    tc_w.set(qn('w:w'), str(round(cm * 567)))
    tc_w.set(qn('w:type'), 'dxa')


def font_run(run, name='Times New Roman', size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:cs'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def format_para(p, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, line=1.15,
                left=0, first=0, keep=False):
    p.alignment = align
    f = p.paragraph_format
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.line_spacing = line
    f.left_indent = Cm(left)
    f.first_line_indent = Cm(first)
    f.keep_with_next = keep
    return p


def add_paragraph(text='', *, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                  before=0, after=6, line=1.15, first=0.75, size=12, keep=False):
    p = doc.add_paragraph()
    format_para(p, align=align, before=before, after=after, line=line, first=first, keep=keep)
    r = p.add_run(text)
    font_run(r, size=size, bold=bold, italic=italic)
    return p


def add_heading(text, page_break=False):
    p = doc.add_paragraph()
    format_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=6, line=1.0, first=0, keep=True)
    p.paragraph_format.page_break_before = page_break
    r = p.add_run(text)
    font_run(r, size=12, bold=True)
    return p


def add_subheading(text):
    p = doc.add_paragraph()
    format_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=4, line=1.0, first=0, keep=True)
    r = p.add_run(text)
    font_run(r, size=12, bold=True)
    return p


def add_number(text, number):
    p = doc.add_paragraph()
    format_para(p, left=0.75, first=-0.75, after=5, line=1.15)
    r = p.add_run(f'{number}. ')
    font_run(r, bold=True)
    r2 = p.add_run(text)
    font_run(r2)
    return p


def add_bullet(text):
    p = doc.add_paragraph()
    format_para(p, left=0.75, first=-0.5, after=4, line=1.15)
    r = p.add_run('• ')
    font_run(r)
    r2 = p.add_run(text)
    font_run(r2)
    return p


def add_code(lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_width(cell, 13.6)
    shade_cell(cell, 'F3F4F6')
    set_cell_border(cell, top={'val':'single','sz':'4','color':'D1D5DB'},
                    bottom={'val':'single','sz':'4','color':'D1D5DB'},
                    left={'val':'single','sz':'4','color':'D1D5DB'},
                    right={'val':'single','sz':'4','color':'D1D5DB'})
    p = cell.paragraphs[0]
    format_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.0, first=0)
    for i, line in enumerate(lines.splitlines()):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        font_run(r, name='Consolas', size=8.5)
    return table


def add_figure(path, caption, width_cm=4.0):
    table = doc.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0,0)
    set_width(cell, 13.6)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell, top={'val':'nil'}, bottom={'val':'nil'}, left={'val':'nil'}, right={'val':'nil'})
    p = cell.paragraphs[0]
    format_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, line=1.0, first=0, keep=True)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    c = table.cell(1,0)
    set_width(c, 13.6)
    set_cell_border(c, top={'val':'nil'}, bottom={'val':'nil'}, left={'val':'nil'}, right={'val':'nil'})
    cp = c.paragraphs[0]
    format_para(cp, align=WD_ALIGN_PARAGRAPH.CENTER, after=5, line=1.0, first=0)
    cr = cp.add_run(caption)
    font_run(cr, size=10, italic=True)
    return table


def add_two_figures(items, width_cm=3.75):
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col, (path, caption) in enumerate(items):
        imgcell = table.cell(0, col)
        capcell = table.cell(1, col)
        set_width(imgcell, 6.8); set_width(capcell, 6.8)
        for cell in (imgcell, capcell):
            set_cell_border(cell, top={'val':'nil'}, bottom={'val':'nil'}, left={'val':'nil'}, right={'val':'nil'})
        p = imgcell.paragraphs[0]
        format_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, line=1.0, first=0, keep=True)
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        cp = capcell.paragraphs[0]
        format_para(cp, align=WD_ALIGN_PARAGRAPH.CENTER, after=4, line=1.0, first=0)
        r = cp.add_run(caption)
        font_run(r, size=9.5, italic=True)
    return table


def replace_para(paragraph, text):
    text_runs = [r for r in paragraph.runs if not r._r.xpath('.//w:drawing')]
    if text_runs:
        text_runs[0].text = text
        for r in text_runs[1:]:
            r.text = ''
    else:
        r = paragraph.add_run(text)
        font_run(r)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('Halaman ')
    font_run(run, size=10)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(fld)

# Cover: preserve drawings and original layout.
replace_para(doc.paragraphs[1], 'MODUL XII – AKSES LOKASI DENGAN GPS DI FLUTTER')
replace_para(doc.paragraphs[14], '\tTanggal\t: 22 Oktober 2026')

# Remove template instruction body, retaining terminal sectPr.
body = doc._body._body
last_cover = doc.paragraphs[22]._p
seen = False
for child in list(body):
    if child is last_cover:
        seen = True
        continue
    if seen and child.tag != qn('w:sectPr'):
        body.remove(child)

# Start the body at its heading; avoid an overflowing blank break paragraph.

# Global style and footer.
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Times New Roman')
normal._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Times New Roman')
normal._element.get_or_add_rPr().rFonts.set(qn('w:cs'), 'Times New Roman')
normal.font.size = Pt(12)
for section in doc.sections:
    fp = section.footer.paragraphs[0]
    fp.clear()
    add_page_number(fp)

add_heading('I. Tujuan', page_break=True)
add_paragraph('Praktikum ini bertujuan memahami proses pengambilan lokasi perangkat dengan GPS pada Flutter, memperoleh koordinat latitude dan longitude, mengubah koordinat menjadi alamat melalui reverse geocoding, menangani izin lokasi Android, serta menampilkan status lokasi secara dinamis dan responsif pada antarmuka aplikasi.')
for text in [
    'Menggunakan paket geolocator untuk memeriksa layanan lokasi, mengelola izin, dan memperoleh posisi perangkat.',
    'Menggunakan paket geocoding untuk menerjemahkan koordinat menjadi informasi kelurahan/kecamatan dan kota.',
    'Menerapkan StatefulWidget dan setState() agar tampilan berubah sesuai keadaan awal, pemrosesan, berhasil, atau error.',
    'Membangun antarmuka menggunakan Scaffold, Card, CircularProgressIndicator, dan ElevatedButton.',
]: add_bullet(text)

add_heading('II. Langkah Kerja')
add_number('Membuat proyek Flutter baru dengan nama lokasi_flutter dan memeriksa struktur awal proyek, terutama folder lib, android, ios, serta file pubspec.yaml.', 1)
add_number('Menambahkan dependensi google_fonts, geolocator, dan geocoding pada pubspec.yaml, kemudian menjalankan flutter pub get. Versi geolocator dan geocoding dinaikkan ke versi kompatibel toolchain saat ini karena geolocator lama belum menyediakan Android namespace, sedangkan geocoding lama memakai compileSdk 33 yang tidak kompatibel dengan dependensi Android saat ini. API diperbarui menjadi LocationSettings dan Geocoding().placemarkFromCoordinates(); fungsi praktikum tetap sama.', 2)
add_code("dependencies:\n  google_fonts: ^6.1.0\n  geolocator: ^14.0.2\n  geocoding: ^5.0.0")
add_number('Menambahkan izin ACCESS_FINE_LOCATION dan ACCESS_COARSE_LOCATION pada android/app/src/main/AndroidManifest.xml agar aplikasi dapat meminta akses lokasi presisi maupun perkiraan.', 3)
add_code('<uses-permission android:name="android.permission.ACCESS_FINE_LOCATION" />\n<uses-permission android:name="android.permission.ACCESS_COARSE_LOCATION" />')
add_number('Membuat MyApp dengan MaterialApp, tema Poppins, dan LocationScreen sebagai halaman utama. LocationScreen menggunakan StatefulWidget karena data lokasi dan status proses berubah selama aplikasi berjalan.', 4)
add_number('Menyediakan empat variabel state: _isLoading untuk status proses, _kecamatan dan _kota untuk hasil reverse geocoding, serta _errorMessage untuk pesan kegagalan. Nilai-nilai ini menentukan konten yang ditampilkan pada Card.', 5)
add_code("String? _kecamatan;\nString? _kota;\nbool _isLoading = false;\nString? _errorMessage;")
add_number('Mengimplementasikan fungsi _getLocation(). Fungsi mengaktifkan loading dan mereset data lama; memeriksa layanan GPS; memeriksa atau meminta izin; mengambil posisi dengan akurasi tinggi; mengubah koordinat menjadi Placemark; lalu menyimpan kecamatan dan kota. Blok catch menampilkan error, sedangkan finally selalu menghentikan loading.', 6)
add_paragraph('Peran setState() sangat penting karena setiap perubahan state harus diberitahukan kepada Flutter. Panggilan setState() di awal fungsi menampilkan spinner dan menonaktifkan tombol. Setelah proses berhasil, setState() mengisi _kecamatan serta _kota. Jika terjadi pengecualian, setState() mengisi _errorMessage. Pada bagian finally, setState() mengembalikan _isLoading ke false sehingga tombol aktif kembali. Setiap perubahan tersebut memicu build() dan membuat antarmuka otomatis mengikuti data terbaru.', first=0.75)
add_number('Menyusun UI dengan Scaffold, AppBar berwarna indigo, Card sebagai area status, dan ElevatedButton untuk memanggil _getLocation(). Percabangan pada Card menampilkan spinner ketika loading, teks merah ketika error, instruksi pada keadaan awal, atau hasil alamat ketika berhasil.', 7)
add_number('Menjalankan flutter test dan flutter analyze. Dua widget test lulus dan analisis statis tidak menemukan masalah. Selanjutnya aplikasi dibangun menjadi APK debug, dipasang pada emulator Android API 36, lalu diuji menggunakan koordinat emulator di Dinoyo, Kota Malang.', 8)

add_subheading('Hasil pengamatan selama pengujian')
add_two_figures([
    (EVIDENCE/'01_tampilan_awal.png', 'Gambar 1. Kondisi awal sebelum lokasi diminta.'),
    (EVIDENCE/'02_loading.png', 'Gambar 2. Proses pengambilan lokasi; spinner aktif dan tombol dinonaktifkan.'),
])
add_paragraph('Pada kondisi awal, Card menampilkan petunjuk untuk menekan tombol. Sesaat setelah tombol ditekan, _isLoading bernilai true, CircularProgressIndicator ditampilkan, dan tombol tidak dapat ditekan ulang. Hierarki UI Android mengonfirmasi properti tombol berada pada kondisi enabled=false selama proses.', first=0.75)
add_two_figures([
    (EVIDENCE/'03_hasil_lokasi.png', 'Gambar 3. Reverse geocoding berhasil: Dinoyo, Kota Malang.'),
    (EVIDENCE/'04_error_izin.png', 'Gambar 4. Error ketika izin lokasi ditolak permanen.'),
])
add_paragraph('Dengan koordinat uji emulator di wilayah Malang, geocoding menghasilkan Kelurahan/Kecamatan Dinoyo dan Kota Malang. Skenario izin ditolak permanen menghasilkan pesan merah yang memberi petunjuk agar izin diubah pada pengaturan aplikasi. Pada kedua akhir proses, loading berhenti dan tombol kembali aktif.', first=0.75)

add_heading('III. Screenshot Hasil', page_break=True)
add_two_figures([
    (EVIDENCE/'01_tampilan_awal.png', 'Screenshot 1. Tampilan awal aplikasi.'),
    (EVIDENCE/'02_loading.png', 'Screenshot 2. Aplikasi sedang memproses lokasi.'),
], width_cm=4.05)
add_two_figures([
    (EVIDENCE/'03_hasil_lokasi.png', 'Screenshot 3. Lokasi Dinoyo, Kota Malang berhasil ditampilkan.'),
    (EVIDENCE/'04_error_izin.png', 'Screenshot 4. Pesan error izin lokasi permanen.'),
], width_cm=4.05)

add_heading('IV. Kesimpulan', page_break=True)
add_paragraph('Aplikasi lokasi_flutter berhasil menerapkan akses lokasi perangkat menggunakan geolocator dan mengubah koordinat menjadi alamat dengan geocoding. Pengujian nyata pada emulator menunjukkan empat kondisi antarmuka yang diminta: tampilan awal, proses pengambilan lokasi, hasil berhasil berupa Dinoyo–Kota Malang, dan error ketika izin lokasi ditolak permanen.')
add_paragraph('Fungsi _getLocation() menggabungkan pemeriksaan layanan GPS, pemeriksaan dan permintaan izin, pengambilan posisi, reverse geocoding, serta penanganan kesalahan. Variabel _isLoading, _kecamatan, _kota, dan _errorMessage menjadi sumber keadaan untuk menentukan konten Card. Pemanggilan setState() setelah perubahan data menyebabkan Flutter menjalankan ulang build(), sehingga spinner, pesan error, hasil alamat, dan status tombol diperbarui otomatis tanpa membuka halaman baru.')
add_paragraph('Hasil praktikum juga menunjukkan bahwa penanganan izin merupakan bagian penting dari aplikasi berbasis lokasi. Aplikasi tidak berhenti ketika izin atau layanan tidak tersedia, melainkan memberi umpan balik yang dapat dipahami pengguna. Dengan demikian, implementasi telah memenuhi tujuan praktikum dan tugas Modul 12 secara fungsional serta responsif.')

# Enforce body typography without deleting drawing/page-break runs.
for p in doc.paragraphs:
    if p._p.xpath('.//w:drawing'):
        continue
    for run in p.runs:
        if run.text and p._p is not doc.paragraphs[0]._p and p._p is not doc.paragraphs[1]._p:
            if run.font.name != 'Consolas':
                font_run(run, size=run.font.size.pt if run.font.size else 12,
                         bold=bool(run.bold), italic=bool(run.italic))
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    if run.text and run.font.name != 'Consolas':
                        font_run(run, size=run.font.size.pt if run.font.size else 12,
                                 bold=bool(run.bold), italic=bool(run.italic))

add_paragraph('Pengujian dilakukan hanya pada emulator Android dengan koordinat simulasi longitude 112.6078 dan latitude -7.9516; bukan lokasi fisik pengguna. Performa smartphone fisik dan iOS belum diuji. Aplikasi mengambil lokasi sekali setiap tombol ditekan, bukan melacak lokasi terus-menerus. Makna subLocality dan subAdministrativeArea dapat berbeda menurut penyedia alamat atau wilayah. Pengambilan posisi dan geocoding belum diberi timeout khusus.')

# Metadata and save aliases.
doc.core_properties.title = 'Laporan Praktikum Modul 12 – Akses Lokasi dengan GPS di Flutter'
doc.core_properties.author = 'Erai Bagusalim (240605110088)'
doc.core_properties.subject = 'Mobile Programming Practicum A'
doc.save(FORMAL)
copy2(FORMAL, SHORT)
copy2(FORMAL, TEMPLATE)

source = ROOT/'lokasi_flutter'/'lib'/'main.dart'
manifest = {
    'module': 12,
    'student': {'name':'Erai Bagusalim','nim':'240605110088','class':'Mobile Programming Practicum A'},
    'source': str(source.relative_to(ROOT)),
    'source_sha256': sha256(source.read_bytes()).hexdigest(),
    'screenshots': [
        {'figure':1,'file':'evidence/01_tampilan_awal.png','state':'awal','observed':'Instruksi awal dan tombol aktif'},
        {'figure':2,'file':'evidence/02_loading.png','state':'loading','observed':'Spinner tampil; tombol enabled=false pada UI hierarchy'},
        {'figure':3,'file':'evidence/03_hasil_lokasi.png','state':'sukses','observed':'Kelurahan/Kecamatan Dinoyo; Kota Malang'},
        {'figure':4,'file':'evidence/04_error_izin.png','state':'error','observed':'Izin lokasi ditolak permanen'},
    ],
    'outputs': [FORMAL.name, SHORT.name, TEMPLATE.name],
}
(EVIDENCE/'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'formal':str(FORMAL),'short':str(SHORT),'template':str(TEMPLATE),'manifest':str(EVIDENCE/'manifest.json')}, ensure_ascii=False, indent=2))
