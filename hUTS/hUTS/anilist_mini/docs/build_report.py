import os
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client

def set_cell_margins(cell, top=45, bottom=45, left=75, right=75):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top=True, bottom=True, left=True, right=True, color="000000", sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for b_name, active in borders.items():
        node = OxmlElement(f'w:{b_name}')
        if active:
            node.set(qn('w:val'), 'single')
            node.set(qn('w:sz'), sz)
            node.set(qn('w:space'), '0')
            node.set(qn('w:color'), color)
        else:
            node.set(qn('w:val'), 'none')
        tcBorders.append(node)
    tcPr.append(tcBorders)

def set_cell_background(cell, fill_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def build_report():
    doc = docx.Document()

    # Standard A4 margins: 2.54 cm all sides
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    # Base Normal Style
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)

    def add_p(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, space_before=0, space_after=2.5, line_spacing=1.15):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_subheading(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_caption(label, keterangan):
        p1 = doc.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.keep_with_next = True
        r1 = p1.add_run(label)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(12)
        r1.font.bold = True

        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(3)
        r2 = p2.add_run(f"Keterangan: {keterangan}")
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(12)

    def add_snippet_caption(label):
        p1 = doc.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(1.5)
        p1.paragraph_format.space_after = Pt(3)
        r1 = p1.add_run(label)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(12)
        r1.font.bold = True

    # =============================================================
    # PAGE 1: COVER
    # =============================================================
    add_p("LAPORAN UTS MOBILE PROGRAMMING", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_before=15, space_after=2)
    add_p("(APLIKASI ANILIST MINI)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, italic=True, space_before=0, space_after=35)

    if os.path.exists('logo_uin.jpeg'):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(10)
        p_logo.paragraph_format.space_after = Pt(45)
        run_logo = p_logo.add_run()
        run_logo.add_picture('logo_uin.jpeg', width=Cm(6.0))

    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.left_indent = Cm(4.8)
    p_info.paragraph_format.space_before = Pt(10)
    p_info.paragraph_format.space_after = Pt(65)
    p_info.paragraph_format.line_spacing = 1.3

    r = p_info.add_run("Nama   : Erai Bagusalim\n")
    r.font.name = 'Times New Roman'; r.font.size = Pt(12)
    r = p_info.add_run("NIM    : 240605110088\n")
    r.font.name = 'Times New Roman'; r.font.size = Pt(12)
    r = p_info.add_run("Kelas  : A")
    r.font.name = 'Times New Roman'; r.font.size = Pt(12)

    add_p("JURUSAN TEKNIK INFORMATIKA", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_before=15, space_after=2)
    add_p("FAKULTAS SAINS DAN TEKNOLOGI", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_before=0, space_after=2)
    add_p("UNIVERSITAS ISLAM NEGERI MAULANA MALIK IBRAHIM MALANG", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_before=0, space_after=2)
    add_p("2026", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_before=0, space_after=0)

    doc.add_page_break()

    # =============================================================
    # PAGE 2: BAB 1. PENDAHULUAN & 1.1 Latar Belakang
    # =============================================================
    add_heading_1("BAB 1. PENDAHULUAN")
    add_heading_2("1.1 Latar Belakang dan Gambaran Aplikasi")
    
    add_p("Anime adalah media hiburan animasi asal Jepang yang sangat digemari dan ditonton oleh jutaan pengguna setiap harinya. Seiring meluasnya penggunaan telepon pintar, semakin banyak orang yang membaca referensi, mengecek skor rating, dan mencari judul anime langsung dari perangkat genggam mereka. Meskipun demikian, mencari anime tertentu di antara ribuan judul, mengetahui skor penilaian, jumlah episodenya, format tayangannya, serta membaca sinopsis ringkasnya masih memerlukan beberapa langkah jika dilakukan melalui penelusuran peramban yang tampilannya kurang ringkas pada layar ponsel.")

    add_p("Berdasarkan kebutuhan tersebut, dikembangkan aplikasi AniList Mini, yaitu aplikasi mobile berbasis Flutter yang menyajikan katalog anime dan rincian informasinya dalam tampilan yang modern, sederhana, dan mudah dicari. Seluruh data diambil secara dinamis dari REST API publik (Tenrai API v1 yang kompatibel dengan skema Jikan API v4), sehingga aplikasi tidak menyimpan data anime secara tetap di dalam kode program.")

    add_p("Tujuan pengembangan aplikasi ini adalah sebagai berikut.")
    
    p = add_p()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("1. Membangun aplikasi mobile menggunakan Flutter dan bahasa Dart yang menyajikan data dari API.")
    r.font.name = 'Times New Roman'; r.font.size = Pt(12)

    p = add_p()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("2. Menerapkan pengambilan data melalui permintaan HTTP dan mengonversi respons JSON menjadi objek Dart melalui class model AnimeModel.")
    r.font.name = 'Times New Roman'; r.font.size = Pt(12)

    p = add_p()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("3. Menerapkan asynchronous programming menggunakan Future dan FutureBuilder dengan penanganan kondisi loading, success, dan error.")
    r.font.name = 'Times New Roman'; r.font.size = Pt(12)

    p = add_p()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("4. Menghasilkan aplikasi yang bermanfaat bagi pengguna yang ingin mencari dan mengeksplorasi katalog anime beserta informasinya.")
    r.font.name = 'Times New Roman'; r.font.size = Pt(12)

    add_p("Sasaran pengguna aplikasi ini adalah pelajar, mahasiswa, dan masyarakat umum berbahasa Indonesia yang ingin mencari anime, memahami rating dan informasi singkatnya, serta membaca sinopsis cerita melalui telepon pintar.")

    add_p("Secara umum, aplikasi terdiri atas dua halaman. Halaman pertama menampilkan daftar anime terpopuler yang dapat dicari. Halaman kedua menampilkan detail satu anime terpilih, berisi poster besar, rating, tipe, episode, tahun, status, daftar genre, dan sinopsis lengkap. Kedua halaman mengambil datanya dari API dan menampilkan indikator yang sesuai saat data sedang dimuat atau gagal dimuat.")

    doc.add_page_break()

    # =============================================================
    # PAGE 3: 1.2 Fitur Aplikasi & BAB 2. PEMBAHASAN 2.1
    # =============================================================
    add_heading_2("1.2 Fitur Aplikasi")
    add_p("Fitur-fitur utama yang tersedia pada aplikasi:")

    t1_data = [
        ["No", "Fitur", "Fungsi", "Manfaat bagi Pengguna"],
        ["1", "Daftar Top Anime", "Menampilkan 24 anime terpopuler berisi poster, judul, rating, dan tipe.", "Pengguna melihat seluruh anime populer dalam satu daftar kisi yang teratur."],
        ["2", "Pencarian anime", "Menyaring daftar anime dari server berdasarkan kata kunci judul.", "Anime yang dicari dapat ditemukan dengan cepat dan akurat."],
        ["3", "Detail anime", "Menampilkan rincian lengkap: poster, rating, tipe, episode, tahun, status, genre, sinopsis.", "Pengguna dapat membaca dan memahami latar belakang anime sebelum menonton."],
        ["4", "Refresh data", "Memuat ulang katalog Top Anime dan membersihkan kata kunci pencarian.", "Pengguna dapat memperbarui daftar katalog kapan saja secara instan."],
        ["5", "Indikator loading", "Menampilkan indikator saat data sedang diambil dari server.", "Pengguna tahu bahwa aplikasi sedang bekerja mengambil data."],
        ["6", "Penanganan error dan coba lagi", "Menampilkan pesan kesalahan yang jelas dan tombol untuk mengulangi permintaan.", "Pengguna tidak melihat layar kosong ketika koneksi bermasalah."]
    ]

    t1 = doc.add_table(rows=len(t1_data), cols=4)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths_1 = [Cm(0.9), Cm(3.2), Cm(6.1), Cm(5.7)]

    for row_idx, row in enumerate(t1.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths_1[col_idx]
            set_cell_margins(cell, top=50, bottom=50, left=75, right=75)
            set_cell_borders(cell, top=True, bottom=True, left=True, right=True, sz="4")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (row_idx == 0 or col_idx == 0) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(t1_data[row_idx][col_idx])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            if row_idx == 0:
                run.font.bold = True
                set_cell_background(cell, "F0F0F0")

    add_p("", space_before=4, space_after=4)
    add_heading_1("BAB 2. PEMBAHASAN")
    add_heading_2("2.1 Persiapan Project dan Package")
    add_subheading("a. Pembuatan project Flutter")
    add_p("Project dibuat menggunakan Android Studio / Flutter CLI dengan memilih jenis project Application dan memberi nama anilist_mini. Aplikasi dijalankan dan diuji pada emulator Android Pixel 5 yang tersedia di Android Studio serta pada Google Chrome dengan tampilan mobile.")

    add_subheading("b. Struktur folder dan berkas utama")

    doc.add_page_break()

    # =============================================================
    # PAGE 4: Tabel Folder & Tabel Package & 2.2 Desain UI
    # =============================================================
    t2_data = [
        ["Folder / Berkas", "Fungsi"],
        ["lib/main.dart", "Titik masuk aplikasi; mengatur tema Material 3 Dark Mode dan halaman awal."],
        ["lib/models/anime_model.dart", "Class model AnimeModel untuk data pada katalog dan detail anime."],
        ["lib/services/anime_service.dart", "Mengirim permintaan HTTP, memeriksa status respons, mengubah JSON, dan menangani kesalahan."],
        ["lib/pages/home_page.dart", "Halaman daftar anime dengan pencarian dan FutureBuilder."],
        ["lib/pages/detail_page.dart", "Halaman detail anime dengan FutureBuilder untuk rincian metadata lengkap."],
        ["test/anime_model_test.dart", "Unit test untuk memeriksa konversi JSON ke model dan penanganan nilai null."],
        ["test/anime_service_test.dart", "Unit test untuk service: endpoint, query encoding, status 429/500, dan jaringan."],
        ["test/widget_test.dart", "Widget test untuk memverifikasi penampilan kartu anime dan kondisi hasil kosong."],
        ["pubspec.yaml", "Daftar dependensi dan konfigurasi project."]
    ]

    t2 = doc.add_table(rows=len(t2_data), cols=2)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths_2 = [Cm(5.4), Cm(10.5)]

    for row_idx, row in enumerate(t2.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths_2[col_idx]
            set_cell_margins(cell, top=45, bottom=45, left=75, right=75)
            set_cell_borders(cell, top=True, bottom=True, left=True, right=True, sz="4")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(t2_data[row_idx][col_idx])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            if row_idx == 0:
                run.font.bold = True
                set_cell_background(cell, "F0F0F0")

    add_p("", space_before=4, space_after=2)
    add_subheading("c. Package yang digunakan")

    t3_data = [
        ["Package", "Fungsi", "Alasan Penggunaan"],
        ["http", "Mengirim permintaan HTTP GET ke API.", "Package resmi yang sederhana dan direkomendasikan dokumentasi Flutter untuk pengambilan data. Juga menyediakan MockClient untuk pengujian tanpa internet."],
        ["flutter_lints", "Aturan analisis kode (lint).", "Menjaga gaya penulisan kode tetap konsisten dan mendeteksi kesalahan umum."],
        ["flutter_test", "Unit test dan widget test.", "Memeriksa fungsi aplikasi secara sistematis dan otomatis."]
    ]

    t3 = doc.add_table(rows=len(t3_data), cols=3)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths_3 = [Cm(2.8), Cm(5.1), Cm(8.0)]

    for row_idx, row in enumerate(t3.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths_3[col_idx]
            set_cell_margins(cell, top=45, bottom=45, left=75, right=75)
            set_cell_borders(cell, top=True, bottom=True, left=True, right=True, sz="4")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(t3_data[row_idx][col_idx])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            if row_idx == 0:
                run.font.bold = True
                set_cell_background(cell, "F0F0F0")

    add_p("", space_before=4, space_after=2)
    add_heading_2("2.2 Desain Antarmuka Pengguna (User Interface)")
    add_p("Antarmuka dirancang elegan dengan tema Material 3 Dark Mode (latar #10101C dan kontainer kartu #1C1B2B). Aplikasi terdiri atas dua halaman, yaitu halaman daftar anime dan halaman detail anime. Komponen yang digunakan antara lain Scaffold dan AppBar, TextField untuk pencarian, GridView.builder untuk daftar katalog anime yang panjang, Card dan InkWell untuk setiap anime, serta Chip dan Wrap untuk daftar genre anime.")

    doc.add_page_break()

    # =============================================================
    # PAGE 5: a. Halaman daftar anime (Gambar 2.1)
    # =============================================================
    add_subheading("a. Halaman daftar anime")
    add_p("Halaman ini adalah halaman utama. Pada bagian atas terdapat kolom pencarian, dan di bawahnya daftar anime. Setiap baris menampilkan poster, judul anime, skor rating bintang, dan tipe tayangan. Mengetuk salah satu baris membuka halaman detail.")

    p_img1 = doc.add_paragraph()
    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img1.paragraph_format.space_before = Pt(6)
    p_img1.paragraph_format.space_after = Pt(4)
    run_img1 = p_img1.add_run()
    run_img1.add_picture('anilist_mini/docs/screenshots/home.png', width=Cm(6.5))
    add_caption("Gambar 2.1 Halaman Daftar Anime", "Daftar 24 anime terpopuler beserta kolom pencarian pada bagian atas.")

    doc.add_page_break()

    # =============================================================
    # PAGE 6: Gambar 2.2 Pencarian & b. Halaman detail anime
    # =============================================================
    p_img2 = doc.add_paragraph()
    p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img2.paragraph_format.space_before = Pt(6)
    p_img2.paragraph_format.space_after = Pt(4)
    run_img2 = p_img2.add_run()
    run_img2.add_picture('anilist_mini/docs/screenshots/search.png', width=Cm(6.5))
    add_caption("Gambar 2.2 Fitur Pencarian Anime", "Daftar tersaring sesuai kata kunci yang diketik pengguna.")

    add_subheading("b. Halaman detail anime")
    add_p("Bagian atas halaman detail berisi informasi anime (poster besar, judul, rating bintang, tipe tayangan, jumlah episode, tahun rilis, dan status). Di bawahnya terdapat kartu daftar genre dan sinopsis cerita yang dapat digulir.")

    doc.add_page_break()

    # =============================================================
    # PAGE 7: Gambar 2.3 Detail & c. Tampilan loading dan error
    # =============================================================
    p_img3 = doc.add_paragraph()
    p_img3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img3.paragraph_format.space_before = Pt(6)
    p_img3.paragraph_format.space_after = Pt(4)
    run_img3 = p_img3.add_run()
    run_img3.add_picture('anilist_mini/docs/screenshots/detail.png', width=Cm(6.5))
    add_caption("Gambar 2.3 Halaman Detail Anime", "Kartu informasi anime dan sinopsis beserta metadatanya.")

    add_subheading("c. Tampilan loading dan error")
    add_p("Saat data sedang diambil, aplikasi menampilkan indikator berputar disertai teks. Jika pengambilan data gagal, aplikasi menampilkan ikon, pesan kesalahan, dan tombol 'Coba Lagi'.")

    doc.add_page_break()

    # =============================================================
    # PAGE 8: Gambar 2.4 Loading
    # =============================================================
    p_img4 = doc.add_paragraph()
    p_img4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img4.paragraph_format.space_before = Pt(12)
    p_img4.paragraph_format.space_after = Pt(4)
    run_img4 = p_img4.add_run()
    run_img4.add_picture('anilist_mini/docs/screenshots/loading.png', width=Cm(6.5))
    add_caption("Gambar 2.4 Tampilan Kondisi Loading", "Indikator ditampilkan selama aplikasi menunggu respons dari server.")

    doc.add_page_break()

    # =============================================================
    # PAGE 9: Gambar 2.5 Error
    # =============================================================
    p_img5 = doc.add_paragraph()
    p_img5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img5.paragraph_format.space_before = Pt(12)
    p_img5.paragraph_format.space_after = Pt(4)
    run_img5 = p_img5.add_run()
    run_img5.add_picture('anilist_mini/docs/screenshots/error.png', width=Cm(6.5))
    add_caption("Gambar 2.5 Tampilan Kondisi Error", "Pesan kesalahan dan tombol 'Coba Lagi' ditampilkan ketika permintaan gagal.")

    doc.add_page_break()

    # =============================================================
    # PAGE 10: 2.3 Implementasi & URL API & Struktur Respons JSON & Tabel Atribut
    # =============================================================
    add_heading_2("2.3 Implementasi Pengambilan dan Pengolahan Data dari API")
    add_subheading("a. URL API dan jenis data")
    add_p("Aplikasi menggunakan API publik Tenrai versi 1 (kompatibel skema Jikan API v4) dengan tiga endpoint berikut.")

    t4_data = [
        ["Endpoint", "Metode", "Data yang Diberikan"],
        ["https://api.tenrai.org/v1/top/anime?limit=24", "GET", "Daftar 24 anime terpopuler beserta informasi ringkasnya."],
        ["https://api.tenrai.org/v1/anime?q={query}&limit=24", "GET", "Daftar anime hasil pencarian berdasarkan kata kunci query."],
        ["https://api.tenrai.org/v1/anime/<id>", "GET", "Detail satu anime spesifik beserta seluruh metadatanya."]
    ]

    t4 = doc.add_table(rows=len(t4_data), cols=3)
    t4.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths_4 = [Cm(6.8), Cm(2.1), Cm(7.0)]

    for row_idx, row in enumerate(t4.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths_4[col_idx]
            set_cell_margins(cell, top=35, bottom=35, left=70, right=70)
            set_cell_borders(cell, top=True, bottom=True, left=True, right=True, sz="4")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (row_idx == 0 or col_idx == 1) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(t4_data[row_idx][col_idx])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            if row_idx == 0:
                run.font.bold = True
                set_cell_background(cell, "F0F0F0")

    add_p("", space_before=2, space_after=1)
    add_subheading("b. Struktur respons API")
    add_p("Respons endpoint berbentuk objek JSON dengan daftar anime berada di dalam field 'data':")

    p_code_json = doc.add_paragraph()
    p_code_json.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_code_json.paragraph_format.space_before = Pt(1)
    p_code_json.paragraph_format.space_after = Pt(1)
    run_code_json = p_code_json.add_run()
    run_code_json.add_picture('anilist_mini/docs/snippets/code_json.png', width=Cm(11.5))

    add_p("Atribut yang dipakai aplikasi dari setiap anime yaitu:")

    t5_data = [
        ["Atribut", "Tipe", "Keterangan"],
        ["mal_id", "int", "Nomor urut/identitas unik anime pada basis data MyAnimeList."],
        ["title", "String", "Judul anime dalam huruf Latin."],
        ["images.jpg.image_url", "String", "URL berkas poster visual anime."],
        ["score", "double", "Nilai evaluasi rating rata-rata pengguna."],
        ["episodes", "int", "Jumlah total episode penayangan anime."],
        ["status", "String", "Status rilis penayangan (Finished / Airing)."],
        ["genres", "List<String>", "Daftar kategori genre anime."],
        ["synopsis", "String", "Penjelasan ringkas sinopsis alur cerita anime."]
    ]

    t5 = doc.add_table(rows=len(t5_data), cols=3)
    t5.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths_5 = [Cm(4.5), Cm(2.8), Cm(8.6)]

    for row_idx, row in enumerate(t5.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths_5[col_idx]
            set_cell_margins(cell, top=30, bottom=30, left=70, right=70)
            set_cell_borders(cell, top=True, bottom=True, left=True, right=True, sz="4")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (row_idx == 0 or col_idx == 1) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(t5_data[row_idx][col_idx])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            if row_idx == 0:
                run.font.bold = True
                set_cell_background(cell, "F0F0F0")

    doc.add_page_break()

    # =============================================================
    # PAGE 11: c. Pengiriman HTTP & Potongan Kode 2.1 & d. JSON
    # =============================================================
    add_subheading("c. Proses pengiriman permintaan HTTP")
    add_p("Pengiriman permintaan dilakukan pada class AnimeService. Alamat dasar API disimpan sebagai konstanta, lalu permintaan GET dikirim menggunakan http.Client dengan batas waktu 20 detik. Diterapkan pembatasan jeda permintaan (_waitForRequestSlot) sebesar 400 ms agar terhindar dari rate limit. Setelah respons diterima, status HTTP diperiksa: jika bukan 200, aplikasi melempar AnimeException berisi pesan kesalahan. Jika berhasil, isi respons didekode dengan jsonDecode dan field data diambil.")

    p_c21 = doc.add_paragraph()
    p_c21.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c21.paragraph_format.space_before = Pt(3)
    p_c21.paragraph_format.space_after = Pt(2)
    run_c21 = p_c21.add_run()
    run_c21.add_picture('anilist_mini/docs/snippets/code_2_1.png', width=Cm(12.5))
    add_snippet_caption("Potongan Kode 2.1 Fungsi _get pada anime_service.dart")

    add_subheading("d. Pengolahan JSON menjadi objek Dart")
    add_p("Setiap objek JSON diubah menjadi objek AnimeModel menggunakan factory constructor AnimeModel.fromJson. Nilai angka dikonversi dengan '(json[...] as num?)?.toInt()', sedangkan nilai teks yang mungkin kosong diberi nilai default melalui operator '??'.")

    doc.add_page_break()

    # =============================================================
    # PAGE 12: Potongan Kode 2.2, 2.3, 2.4 & e. Asynchronous programming
    # =============================================================
    p_c22 = doc.add_paragraph()
    p_c22.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c22.paragraph_format.space_before = Pt(2)
    p_c22.paragraph_format.space_after = Pt(2)
    run_c22 = p_c22.add_run()
    run_c22.add_picture('anilist_mini/docs/snippets/code_2_2.png', width=Cm(10.5))
    add_snippet_caption("Potongan Kode 2.2 Factory AnimeModel.fromJson pada anime_model.dart")

    p_c23 = doc.add_paragraph()
    p_c23.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c23.paragraph_format.space_before = Pt(2)
    p_c23.paragraph_format.space_after = Pt(2)
    run_c23 = p_c23.add_run()
    run_c23.add_picture('anilist_mini/docs/snippets/code_2_3.png', width=Cm(11.0))
    add_snippet_caption("Potongan Kode 2.3 Fungsi fetchTopAnime pada anime_service.dart")

    p_c24 = doc.add_paragraph()
    p_c24.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c24.paragraph_format.space_before = Pt(2)
    p_c24.paragraph_format.space_after = Pt(2)
    run_c24 = p_c24.add_run()
    run_c24.add_picture('anilist_mini/docs/snippets/code_2_4.png', width=Cm(11.0))
    add_snippet_caption("Potongan Kode 2.4 Fungsi searchAnime pada anime_service.dart")

    add_subheading("e. Implementasi asynchronous programming")
    add_p("Pengambilan data membutuhkan waktu karena aplikasi berkomunikasi dengan server, sehingga fungsi dibuat asinkron dengan kata kunci async dan await dan mengembalikan Future.")

    doc.add_page_break()

    # =============================================================
    # PAGE 13: Potongan Kode 2.5 & f. Penanganan Loading, Success, Error
    # =============================================================
    add_p("Pada halaman daftar, Future disimpan pada variabel _animeFuture di dalam initState. Dengan cara ini permintaan hanya dikirim satu kali saat halaman dibuat, dan tidak diulang setiap kali widget dibangun ulang.")

    p_c25 = doc.add_paragraph()
    p_c25.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c25.paragraph_format.space_before = Pt(3)
    p_c25.paragraph_format.space_after = Pt(2)
    run_c25 = p_c25.add_run()
    run_c25.add_picture('anilist_mini/docs/snippets/code_2_5.png', width=Cm(10.5))
    add_snippet_caption("Potongan Kode 2.5 Inisialisasi Future pada initState")

    add_subheading("f. Penanganan kondisi loading, success, dan error")
    add_p("Tampilan halaman disesuaikan dengan status Future menggunakan FutureBuilder.")

    t6_data = [
        ["Kondisi", "Pemicu", "Tampilan"],
        ["Loading", "connectionState belum done (menunggu server).", "Indikator CircularProgressIndicator di tengah layar."],
        ["Error", "snapshot.hasError bernilai true (permintaan gagal).", "Ikon, pesan kesalahan, dan tombol 'Coba Lagi' yang mengulangi permintaan."],
        ["Empty", "snapshot.hasData true dan list kosong.", "Teks 'Anime tidak ditemukan' dan tombol kembali ke Top Anime."],
        ["Success", "Future selesai dan snapshot.data tersedia.", "Kolom pencarian dan katalog anime dalam bentuk kisi (grid)."]
    ]

    t6 = doc.add_table(rows=len(t6_data), cols=3)
    t6.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths_6 = [Cm(2.6), Cm(5.6), Cm(7.7)]

    for row_idx, row in enumerate(t6.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths_6[col_idx]
            set_cell_margins(cell, top=45, bottom=45, left=75, right=75)
            set_cell_borders(cell, top=True, bottom=True, left=True, right=True, sz="4")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (row_idx == 0 or col_idx == 0) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(t6_data[row_idx][col_idx])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            if row_idx == 0:
                run.font.bold = True
                set_cell_background(cell, "F0F0F0")

    doc.add_page_break()

    # =============================================================
    # PAGE 14: Potongan Kode 2.6 & Error Table & g. Penyajian & Potongan 2.7
    # =============================================================
    p_c26 = doc.add_paragraph()
    p_c26.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c26.paragraph_format.space_before = Pt(0)
    p_c26.paragraph_format.space_after = Pt(1)
    run_c26 = p_c26.add_run()
    run_c26.add_picture('anilist_mini/docs/snippets/code_2_6.png', width=Cm(8.4))
    add_snippet_caption("Potongan Kode 2.6 FutureBuilder pada home_page.dart")

    add_p("Jenis kesalahan yang ditangani pada AnimeService beserta pesan yang ditampilkan kepada pengguna:", space_after=1.5)

    t7_data = [
        ["Jenis Kesalahan", "Pesan untuk Pengguna"],
        ["Tidak ada internet (SocketException / ClientException)", "Tidak dapat terhubung ke server. Periksa koneksi internet Anda."],
        ["Waktu habis (TimeoutException, 20 detik)", "Permintaan terlalu lama. Periksa koneksi internet Anda."],
        ["Rate limit (HTTP status 429)", "Terlalu banyak permintaan ke server. Tunggu sebentar, lalu coba lagi."],
        ["Status HTTP bukan 200", "Server mengembalikan kesalahan (HTTP {kode})."],
        ["JSON tidak valid (FormatException)", "Data dari server tidak dapat dibaca."],
        ["Struktur atau tipe data tidak sesuai", "Format data dari server tidak sesuai."]
    ]

    t7 = doc.add_table(rows=len(t7_data), cols=2)
    t7.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths_7 = [Cm(6.8), Cm(9.1)]

    for row_idx, row in enumerate(t7.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths_7[col_idx]
            set_cell_margins(cell, top=25, bottom=25, left=65, right=65)
            set_cell_borders(cell, top=True, bottom=True, left=True, right=True, sz="4")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(t7_data[row_idx][col_idx])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            if row_idx == 0:
                run.font.bold = True
                set_cell_background(cell, "F0F0F0")

    add_p("", space_before=1, space_after=1)
    add_subheading("g. Penyajian data pada antarmuka")
    add_p("Ketika data berhasil diperoleh, daftar ditampilkan dengan GridView.builder yang hanya membangun item yang terlihat di layar sehingga efisien. Fitur pencarian bekerja dengan menyaring kata kunci tanpa permintaan baru ke server pada setiap huruf:", space_after=1)

    p_c27 = doc.add_paragraph()
    p_c27.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c27.paragraph_format.space_before = Pt(0)
    p_c27.paragraph_format.space_after = Pt(1)
    run_c27 = p_c27.add_run()
    run_c27.add_picture('anilist_mini/docs/snippets/code_2_7.png', width=Cm(8.8))
    add_snippet_caption("Potongan Kode 2.7 Fungsi _search pada home_page.dart")

    add_p("Halaman detail bekerja dengan pola yang sama: DetailPage menerima data anime awal kemudian FutureBuilder memuat informasi detail anime secara lengkap.", space_before=1, space_after=1)

    doc.add_page_break()

    # =============================================================
    # PAGE 15: 2.4 Pengujian Aplikasi & a. Pengujian Otomatis
    # =============================================================
    add_heading_2("2.4 Pengujian Aplikasi")
    add_p("Pengujian dilakukan dengan dua cara: pengujian otomatis menggunakan flutter test (tanpa memerlukan internet karena memakai MockClient) dan pengujian manual pada emulator atau peramban.")

    add_subheading("a. Pengujian otomatis")
    add_p("Perintah yang dijalankan: flutter test.")

    t8_data = [
        ["No", "Pengujian", "Hasil yang Diharapkan", "Hasil"],
        ["1", "AnimeModel.fromJson", "JSON berubah menjadi objek AnimeModel dengan benar.", "Lulus"],
        ["2", "AnimeModel field null", "Nilai default aman diberikan saat field opsional null.", "Lulus"],
        ["3", "fetchTopAnime sukses", "Mengembalikan list AnimeModel sesuai jumlah data.", "Lulus"],
        ["4", "searchAnime encoding", "Query URL terkodekan dan menerima hasil kosong.", "Lulus"],
        ["5", "fetchAnimeDetail sukses", "Mengembalikan objek AnimeModel detail tunggal.", "Lulus"],
        ["6", "Respons status 429", "Melempar AnimeException dengan pesan rate limit.", "Lulus"],
        ["7", "Respons status 500", "Melempar AnimeException dengan status server.", "Lulus"],
        ["8", "Respons bukan JSON", "Melempar AnimeException.", "Lulus"],
        ["9", "Kegagalan jaringan", "Melempar AnimeException pesan koneksi internet.", "Lulus"],
        ["10", "Widget: loading & sukses", "Indikator loading tampil, lalu kartu anime tampil.", "Lulus"]
    ]

    t8 = doc.add_table(rows=len(t8_data), cols=4)
    t8.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths_8 = [Cm(0.9), Cm(5.4), Cm(7.7), Cm(1.9)]

    for row_idx, row in enumerate(t8.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths_8[col_idx]
            set_cell_margins(cell, top=35, bottom=35, left=70, right=70)
            set_cell_borders(cell, top=True, bottom=True, left=True, right=True, sz="4")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (row_idx == 0 or col_idx in [0, 3]) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(t8_data[row_idx][col_idx])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            if row_idx == 0:
                run.font.bold = True
                set_cell_background(cell, "F0F0F0")

    add_p("", space_before=2, space_after=1)

    p_test_img = doc.add_paragraph()
    p_test_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_test_img.paragraph_format.space_before = Pt(2)
    p_test_img.paragraph_format.space_after = Pt(1)
    run_test_img = p_test_img.add_run()
    run_test_img.add_picture('anilist_mini/docs/snippets/flutter_test_output.png', width=Cm(12.5))
    add_caption("Gambar 2.6 Hasil Perintah flutter test", "Seluruh pengujian otomatis beserta statusnya.")

    doc.add_page_break()

    # =============================================================
    # PAGE 16: b. Pengujian manual (Baris 1-3 dengan gambar)
    # =============================================================
    add_subheading("b. Pengujian manual")

    t9_header = ["No", "Skenario", "Langkah", "Hasil yang diharapkan", "Hasil aktual"]
    t9_rows_p16 = [
        ("1", "Memuat daftar anime", "Buka aplikasi dengan internet aktif.", "Loading tampil, lalu 24 anime terpopuler tampil.", "anilist_mini/docs/screenshots/home.png"),
        ("2", "Pencarian ditemukan", "Ketik \"Naruto\" pada kolom pencarian.", "Hanya judul terkait Naruto yang tampil.", "anilist_mini/docs/screenshots/search.png"),
        ("3", "Pencarian tidak ditemukan", "Ketik kata yang tidak ada, misalnya \"xyz\".", "Teks \"Anime tidak ditemukan\" tampil.", "anilist_mini/docs/screenshots/empty.png")
    ]

    t_man_1 = doc.add_table(rows=4, cols=5)
    t_man_1.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w_man = [Cm(0.9), Cm(2.7), Cm(3.4), Cm(4.5), Cm(4.4)]

    # Header
    for c_idx, cell in enumerate(t_man_1.rows[0].cells):
        cell.width = col_w_man[c_idx]
        set_cell_margins(cell, top=40, bottom=40, left=65, right=65)
        set_cell_borders(cell, top=True, bottom=True, left=True, right=True, sz="4")
        set_cell_background(cell, "F0F0F0")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t9_header[c_idx]); r.font.name = 'Times New Roman'; r.font.size = Pt(12); r.font.bold = True

    for r_idx, (no, skenario, langkah, hasil_harap, img_path) in enumerate(t9_rows_p16, start=1):
        row = t_man_1.rows[r_idx]
        for c_idx in range(5):
            cell = row.cells[c_idx]
            cell.width = col_w_man[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=40, bottom=40, left=65, right=65)
            set_cell_borders(cell, top=True, bottom=True, left=True, right=True, sz="4")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(no); r.font.name = 'Times New Roman'; r.font.size = Pt(12)
            elif c_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(skenario); r.font.name = 'Times New Roman'; r.font.size = Pt(12)
            elif c_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(langkah); r.font.name = 'Times New Roman'; r.font.size = Pt(12)
            elif c_idx == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(hasil_harap); r.font.name = 'Times New Roman'; r.font.size = Pt(12)
            elif c_idx == 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if os.path.exists(img_path):
                    r = p.add_run()
                    r.add_picture(img_path, width=Cm(2.1))

    doc.add_page_break()

    # =============================================================
    # PAGE 17: Pengujian manual (Baris 4-6 dengan gambar)
    # =============================================================
    t9_rows_p17 = [
        ("4", "Membuka detail anime", "Ketuk salah satu anime.", "Loading, lalu rincian dan sinopsis tampil.", "anilist_mini/docs/screenshots/detail.png"),
        ("5", "Tanpa internet (daftar)", "Matikan internet, lalu buka aplikasi.", "Pesan error dan tombol \"Coba Lagi\" tampil.", "anilist_mini/docs/screenshots/error.png"),
        ("6", "Coba lagi", "Nyalakan internet, tekan \"Coba Lagi\".", "Loading, lalu daftar anime tampil.", "anilist_mini/docs/screenshots/home.png")
    ]

    t_man_2 = doc.add_table(rows=4, cols=5)
    t_man_2.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header
    for c_idx, cell in enumerate(t_man_2.rows[0].cells):
        cell.width = col_w_man[c_idx]
        set_cell_margins(cell, top=40, bottom=40, left=65, right=65)
        set_cell_borders(cell, top=True, bottom=True, left=True, right=True, sz="4")
        set_cell_background(cell, "F0F0F0")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t9_header[c_idx]); r.font.name = 'Times New Roman'; r.font.size = Pt(12); r.font.bold = True

    for r_idx, (no, skenario, langkah, hasil_harap, img_path) in enumerate(t9_rows_p17, start=1):
        row = t_man_2.rows[r_idx]
        for c_idx in range(5):
            cell = row.cells[c_idx]
            cell.width = col_w_man[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=40, bottom=40, left=65, right=65)
            set_cell_borders(cell, top=True, bottom=True, left=True, right=True, sz="4")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(no); r.font.name = 'Times New Roman'; r.font.size = Pt(12)
            elif c_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(skenario); r.font.name = 'Times New Roman'; r.font.size = Pt(12)
            elif c_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(langkah); r.font.name = 'Times New Roman'; r.font.size = Pt(12)
            elif c_idx == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(hasil_harap); r.font.name = 'Times New Roman'; r.font.size = Pt(12)
            elif c_idx == 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if os.path.exists(img_path):
                    r = p.add_run()
                    r.add_picture(img_path, width=Cm(2.1))

    doc.add_page_break()

    # =============================================================
    # PAGE 18: Pengujian manual (Baris 7-8) & BAB III. PENUTUP
    # =============================================================
    t9_rows_p18 = [
        ("7", "Tanpa internet (detail)", "Matikan internet, ketuk anime.", "Pesan error pada halaman detail tampil.", "anilist_mini/docs/screenshots/error.png"),
        ("8", "Deskripsi anime", "Buka detail anime pada halaman detail.", "Sinopsis dan genre tampil rapi.", "anilist_mini/docs/screenshots/detail.png")
    ]

    t_man_3 = doc.add_table(rows=3, cols=5)
    t_man_3.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header
    for c_idx, cell in enumerate(t_man_3.rows[0].cells):
        cell.width = col_w_man[c_idx]
        set_cell_margins(cell, top=40, bottom=40, left=65, right=65)
        set_cell_borders(cell, top=True, bottom=True, left=True, right=True, sz="4")
        set_cell_background(cell, "F0F0F0")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t9_header[c_idx]); r.font.name = 'Times New Roman'; r.font.size = Pt(12); r.font.bold = True

    for r_idx, (no, skenario, langkah, hasil_harap, img_path) in enumerate(t9_rows_p18, start=1):
        row = t_man_3.rows[r_idx]
        for c_idx in range(5):
            cell = row.cells[c_idx]
            cell.width = col_w_man[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=40, bottom=40, left=65, right=65)
            set_cell_borders(cell, top=True, bottom=True, left=True, right=True, sz="4")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(no); r.font.name = 'Times New Roman'; r.font.size = Pt(12)
            elif c_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(skenario); r.font.name = 'Times New Roman'; r.font.size = Pt(12)
            elif c_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(langkah); r.font.name = 'Times New Roman'; r.font.size = Pt(12)
            elif c_idx == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(hasil_harap); r.font.name = 'Times New Roman'; r.font.size = Pt(12)
            elif c_idx == 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if os.path.exists(img_path):
                    r = p.add_run()
                    r.add_picture(img_path, width=Cm(2.1))

    add_p("", space_before=5, space_after=3)
    add_heading_1("BAB III. PENUTUP")
    add_heading_2("3.1 Kesimpulan")
    add_p("Aplikasi AniList Mini berhasil dikembangkan menggunakan Flutter dan Dart. Aplikasi mengambil data secara dinamis dari API Tenrai melalui permintaan HTTP, memeriksa status respons, mengubah JSON menjadi objek Dart dengan class model AnimeModel, lalu menyajikannya pada dua halaman, yaitu daftar anime dan detail anime.")

    add_p("Asynchronous programming diterapkan menggunakan Future, async/await, dan FutureBuilder sehingga aplikasi dapat menampilkan kondisi loading, success, dan error dengan jelas. Kesalahan seperti tidak ada internet, waktu habis, status server yang salah, dan data yang tidak valid ditangani dengan pesan yang mudah dipahami serta tombol untuk")

    doc.add_page_break()

    # =============================================================
    # PAGE 19: Lanjutan Kesimpulan, 3.2 Saran & DAFTAR PUSTAKA
    # =============================================================
    add_p("Berdasarkan hasil pengujian otomatis dan manual pada Bab 2.4, seluruh skenario berjalan sesuai harapan sehingga aplikasi dinyatakan telah berfungsi sesuai dengan tujuan pengembangannya.")

    add_heading_2("3.2 Saran")
    add_p("Beberapa saran untuk pengembangan pada tahap berikutnya adalah sebagai berikut.")

    saran_items = [
        "Menambahkan penyimpanan lokal (cache) agar daftar anime dan detail yang sudah dibuka tetap dapat dibaca tanpa internet.",
        "Menambahkan fitur penanda (bookmark) dan riwayat terakhir dibuka.",
        "Menerapkan pagination agar aplikasi dapat memuat lebih banyak judul anime.",
        "Menambahkan filter genre dan penyortiran berdasarkan rating atau tahun rilis.",
        "Menerapkan state management seperti Provider atau Riverpod bila fitur bertambah banyak."
    ]

    for s in saran_items:
        p_s = add_p()
        p_s.paragraph_format.left_indent = Cm(0.8)
        p_s.paragraph_format.space_after = Pt(2)
        r = p_s.add_run(f"• {s}")
        r.font.name = 'Times New Roman'; r.font.size = Pt(12)

    add_p("", space_before=4, space_after=2)
    add_heading_1("DAFTAR PUSTAKA")

    dp_list = [
        "Flutter Team. (2026). Flutter documentation. https://docs.flutter.dev (diakses 20 Oktober 2026).",
        "Flutter Team. (2026). Fetch data from the internet. https://docs.flutter.dev/cookbook/networking/fetch-data (diakses 20 Oktober 2026).",
        "Flutter Team. (2026). FutureBuilder class. https://api.flutter.dev/flutter/widgets/FutureBuilder-class.html (diakses 20 Oktober 2026).",
        "Dart Team. (2026). Dart programming language documentation. https://dart.dev/guides (diakses 20 Oktober 2026).",
        "Dart Team. (2026). http package. https://pub.dev/packages/http (diakses 20 Oktober 2026).",
        "Flutter Team. (2026). flutter_lints package. https://pub.dev/packages/flutter_lints (diakses 20 Oktober 2026).",
        "Tenrai. (2026). Tenrai API documentation. https://api.tenrai.org (diakses 20 Oktober 2026)."
    ]

    for dp in dp_list:
        p_dp = doc.add_paragraph()
        p_dp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_dp.paragraph_format.left_indent = Cm(0.8)
        p_dp.paragraph_format.first_line_indent = Cm(-0.8)
        p_dp.paragraph_format.space_before = Pt(1)
        p_dp.paragraph_format.space_after = Pt(2.5)
        p_dp.paragraph_format.line_spacing = 1.15
        r_dp = p_dp.add_run(dp)
        r_dp.font.name = 'Times New Roman'
        r_dp.font.size = Pt(12)

    output_docx = "LAPORAN_UTS_ANILIST_MINI.docx"
    doc.save(output_docx)
    print(f"Report saved to {output_docx}")

    print("Converting DOCX to PDF via MS Word 2021...")
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    in_file = os.path.abspath(output_docx)
    out_pdf = os.path.abspath("LAPORAN_UTS_ANILIST_MINI.pdf")
    wb = word.Documents.Open(in_file)
    wb.SaveAs(out_pdf, FileFormat=17) # 17 = wdFormatPDF
    wb.Close()
    word.Quit()
    print(f"PDF created successfully at: {out_pdf}")

if __name__ == "__main__":
    build_report()
